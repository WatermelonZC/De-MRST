"""Deterministic physical-unit Marsupial instance generation and serialization."""

import json
import math
import random
from datetime import datetime
from dataclasses import asdict, dataclass
from pathlib import Path
from typing import Any, Dict, Mapping, Optional, Sequence, Tuple

from .core import (
    DEFAULT_DETACH_TIME,
    DEFAULT_DOR_SPEED,
    DEFAULT_DOCK_TIME,
    DEFAULT_MBR_SPEED,
    Position,
    TaskSpec,
    validate_role_speed_order,
)


SUPPORTED_DISTRIBUTIONS = ("uniform", "gaussian_mixture", "spiral")
DUE_TIME_PROFILES = (
    "legacy",
    "linear_50_noise10",
    "linear_40_noise10",
    "linear_100_noise10",
    "linear_60_noise15",
    "linear_50_noise15",
    "linear_40_noise15",
)


def generation_kwargs_from_setting(setting: Mapping[str, Any]) -> Dict[str, Any]:
    """Extract physical instance-generation fields from a frozen setting.

    Formal protocol settings are also stored in checkpoint metadata, so keep the
    values JSON-friendly there and normalize only the two range fields at the
    generator boundary.
    """
    keys = (
        "speed",
        "mbr_speed",
        "dor_speed",
        "pickup_time_range",
        "handling_time_range",
        "due_time_profile",
    )
    kwargs = {
        key: setting[key]
        for key in keys
        if key in setting and setting[key] is not None
    }
    for key in ("pickup_time_range", "handling_time_range"):
        if key in kwargs:
            value = kwargs[key]
            if not isinstance(value, (list, tuple)) or len(value) != 2:
                raise ValueError(f"{key} must be a two-value range")
            kwargs[key] = (float(value[0]), float(value[1]))
    return kwargs


@dataclass(frozen=True)
class MarsupialInstance:
    instance_id: str
    distribution: str
    data_seed: int
    tasks: Tuple[TaskSpec, ...]
    n_mbr: int
    n_dor: int
    speed: float = 1.8
    dock_time: float = DEFAULT_DOCK_TIME
    detach_time: float = DEFAULT_DETACH_TIME
    # The physical protocol is heterogeneous: every mother uses 1.2 and every
    # child uses 2.4.  ``speed`` remains the scalar objective normalization
    # retained by the legacy reporting format.
    mbr_speed: float = DEFAULT_MBR_SPEED
    dor_speed: float = DEFAULT_DOR_SPEED
    mbr_initial_positions: Tuple[Position, ...] = ()
    dor_initial_positions: Tuple[Position, ...] = ()

    def __post_init__(self) -> None:
        if self.speed <= 0:
            raise ValueError("speed must be positive")
        for name, value in (("mbr_speed", self.mbr_speed), ("dor_speed", self.dor_speed)):
            if value is None or value <= 0:
                raise ValueError(f"{name} must be positive")
        validate_role_speed_order(
            self.resolved_mbr_speed,
            self.resolved_dor_speed,
            require_strict=True,
        )

        if not self.mbr_initial_positions or not self.dor_initial_positions:
            mbr_positions, dor_positions = _generate_initial_positions(
                self.data_seed, self.n_mbr, self.n_dor
            )
            if not self.mbr_initial_positions:
                object.__setattr__(self, "mbr_initial_positions", mbr_positions)
            if not self.dor_initial_positions:
                object.__setattr__(self, "dor_initial_positions", dor_positions)

        def normalize_positions(
            positions: Sequence[Position], count: int, name: str
        ) -> Tuple[Position, ...]:
            values = tuple((float(item[0]), float(item[1])) for item in positions)
            if len(values) != count:
                raise ValueError(f"{name} initial positions must match fleet size")
            if any(
                not all(math.isfinite(coordinate) for coordinate in position)
                for position in values
            ):
                raise ValueError(f"{name} initial positions must be finite")
            return values

        mbr_positions = normalize_positions(
            self.mbr_initial_positions, self.n_mbr, "MBR"
        )
        dor_positions = normalize_positions(
            self.dor_initial_positions, self.n_dor, "DOR"
        )
        if len(set(mbr_positions + dor_positions)) != self.n_mbr + self.n_dor:
            raise ValueError("all MBR/DOR initial positions must be distinct")
        object.__setattr__(self, "mbr_initial_positions", mbr_positions)
        object.__setattr__(self, "dor_initial_positions", dor_positions)

    @property
    def resolved_mbr_speed(self) -> float:
        return float(self.mbr_speed)

    @property
    def resolved_dor_speed(self) -> float:
        return float(self.dor_speed)

    def to_dict(self) -> Dict[str, object]:
        payload = {
            "schema_version": 1,
            "instance_id": self.instance_id,
            "distribution": self.distribution,
            "data_seed": self.data_seed,
            "n_mbr": self.n_mbr,
            "n_dor": self.n_dor,
            "speed": self.speed,
            "dock_time": self.dock_time,
            "detach_time": self.detach_time,
            "tasks": [asdict(task) for task in self.tasks],
        }
        payload["mbr_speed"] = self.mbr_speed
        payload["dor_speed"] = self.dor_speed
        payload["mbr_initial_positions"] = [list(position) for position in self.mbr_initial_positions]
        payload["dor_initial_positions"] = [list(position) for position in self.dor_initial_positions]
        return payload

    @classmethod
    def from_dict(cls, payload: Dict[str, object]) -> "MarsupialInstance":
        if payload.get("schema_version") != 1:
            raise ValueError("unsupported instance schema")
        tasks = tuple(
            TaskSpec(
                source=tuple(row["source"]),
                destination=tuple(row["destination"]),
                due_time=float(row["due_time"]),
                pickup_time=float(row["pickup_time"]),
                handling_time=float(row["handling_time"]),
            )
            for row in payload["tasks"]
        )
        return cls(
            instance_id=str(payload["instance_id"]),
            distribution=str(payload["distribution"]),
            data_seed=int(payload["data_seed"]),
            tasks=tasks,
            n_mbr=int(payload["n_mbr"]),
            n_dor=int(payload["n_dor"]),
            speed=float(payload["speed"]),
            dock_time=float(payload["dock_time"]),
            detach_time=float(payload["detach_time"]),
            mbr_speed=float(
                DEFAULT_MBR_SPEED
                if payload.get("mbr_speed") is None
                else payload["mbr_speed"]
            ),
            dor_speed=float(
                DEFAULT_DOR_SPEED
                if payload.get("dor_speed") is None
                else payload["dor_speed"]
            ),
            mbr_initial_positions=tuple(
                tuple(position)
                for position in (payload.get("mbr_initial_positions") or ())
            ),
            dor_initial_positions=tuple(
                tuple(position)
                for position in (payload.get("dor_initial_positions") or ())
            ),
        )


def _uniform_point(rng: random.Random) -> Tuple[float, float]:
    return tuple(5.0 * round(rng.uniform(0.0, 100.0) / 5.0) for _ in range(2))


def _gaussian_mixture_point(rng: random.Random) -> Tuple[float, float]:
    centers = ((25.0, 75.0), (75.0, 25.0))
    while True:
        center = centers[rng.randrange(2)]
        point = (rng.gauss(center[0], 10.0), rng.gauss(center[1], 10.0))
        if all(0.0 <= coordinate <= 100.0 for coordinate in point):
            return point


def _spiral_point(rng: random.Random) -> Tuple[float, float]:
    theta = math.sqrt(rng.random()) * 5.0 * math.pi
    radius = theta / (5.0 * math.pi)
    x = 50.0 + 45.0 * radius * math.cos(theta) + rng.gauss(0.0, 2.0)
    y = 50.0 + 45.0 * radius * math.sin(theta) + rng.gauss(0.0, 2.0)
    return max(0.0, min(100.0, x)), max(0.0, min(100.0, y))


def _generate_initial_positions(
    seed: int, n_mbr: int, n_dor: int
) -> Tuple[Tuple[Position, ...], Tuple[Position, ...]]:
    """Generate deterministic, distinct fleet starts without changing tasks."""
    rng = random.Random(int(seed) + 0x5EED_2026)
    positions = []
    while len(positions) < n_mbr + n_dor:
        candidate = (
            5.0 * round(rng.uniform(0.0, 100.0) / 5.0),
            5.0 * round(rng.uniform(0.0, 100.0) / 5.0),
        )
        if candidate not in positions:
            positions.append(candidate)
    return tuple(positions[:n_mbr]), tuple(positions[n_mbr:])


def generate_instance(
    task_count: int,
    n_mbr: int,
    n_dor: int,
    distribution: str,
    seed: int,
    *,
    speed: float = 1.8,
    mbr_speed: Optional[float] = DEFAULT_MBR_SPEED,
    dor_speed: Optional[float] = DEFAULT_DOR_SPEED,
    pickup_time_range: Optional[Tuple[float, float]] = None,
    handling_time_range: Optional[Tuple[float, float]] = None,
    due_time_profile: str = "legacy",
    instance_id: str = None,
) -> MarsupialInstance:
    if task_count <= 0 or n_mbr <= 0 or n_dor <= 0:
        raise ValueError("task and fleet sizes must be positive")
    if distribution not in SUPPORTED_DISTRIBUTIONS:
        raise ValueError(f"unsupported distribution: {distribution}")
    if speed <= 0:
        raise ValueError("speed must be positive")
    for name, value in (("mbr_speed", mbr_speed), ("dor_speed", dor_speed)):
        if value is not None and value <= 0:
            raise ValueError(f"{name} must be positive")
    validate_role_speed_order(
        DEFAULT_MBR_SPEED if mbr_speed is None else float(mbr_speed),
        DEFAULT_DOR_SPEED if dor_speed is None else float(dor_speed),
        require_strict=True,
    )

    def validate_range(name: str, value: Optional[Tuple[float, float]]) -> None:
        if value is None:
            return
        if len(value) != 2 or float(value[0]) < 0 or float(value[1]) < float(value[0]):
            raise ValueError(f"{name} must be a non-negative (low, high) range")

    validate_range("pickup_time_range", pickup_time_range)
    validate_range("handling_time_range", handling_time_range)
    due_time_profile = str(due_time_profile).lower()
    if due_time_profile not in DUE_TIME_PROFILES:
        raise ValueError(
            f"unsupported due_time_profile: {due_time_profile}; "
            f"expected one of {DUE_TIME_PROFILES}"
        )
    rng = random.Random(int(seed))
    point_generator = {
        "uniform": _uniform_point,
        "gaussian_mixture": _gaussian_mixture_point,
        "spiral": _spiral_point,
    }[distribution]
    tasks = []
    for index in range(task_count):
        source = point_generator(rng)
        destination = point_generator(rng)
        if due_time_profile == "linear_50_noise10":
            due_time = float(300.0 + 50.0 * index + 10.0 * rng.random())
        elif due_time_profile == "linear_40_noise10":
            due_time = float(200.0 + 40.0 * index + 10.0 * rng.random())
        elif due_time_profile == "linear_100_noise10":
            due_time = float(200.0 + 100.0 * index + 10.0 * rng.random())
        elif due_time_profile == "linear_60_noise15":
            due_time = float(200.0 + 60.0 * index + 15.0 * rng.random())
        elif due_time_profile == "linear_50_noise15":
            due_time = float(300.0 + 50.0 * index + 15.0 * rng.random())
        elif due_time_profile == "linear_40_noise15":
            due_time = float(300.0 + 40.0 * index + 15.0 * rng.random())
        else:
            due_time = float(round(300.0 + 40.0 * index + rng.uniform(-40.0, 40.0)))
        if pickup_time_range is None:
            pickup_time = 30.0
        else:
            pickup_time = float(rng.uniform(*pickup_time_range))
        if handling_time_range is None:
            handling_time = float(60 + 5 * rng.randrange(9))
        else:
            handling_time = float(rng.uniform(*handling_time_range))
        tasks.append(TaskSpec(
            source=source,
            destination=destination,
            due_time=due_time,
            pickup_time=pickup_time,
            handling_time=handling_time,
        ))
    tasks = tuple(tasks)
    mbr_initial_positions, dor_initial_positions = _generate_initial_positions(
        int(seed), n_mbr, n_dor
    )
    return MarsupialInstance(
        instance_id=instance_id or f"n{task_count}-{seed}",
        distribution=distribution,
        data_seed=int(seed),
        tasks=tasks,
        n_mbr=n_mbr,
        n_dor=n_dor,
        speed=float(speed),
        mbr_speed=DEFAULT_MBR_SPEED if mbr_speed is None else float(mbr_speed),
        dor_speed=DEFAULT_DOR_SPEED if dor_speed is None else float(dor_speed),
        mbr_initial_positions=mbr_initial_positions,
        dor_initial_positions=dor_initial_positions,
    )


def write_instance(instance: MarsupialInstance, path: Path) -> None:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_suffix(path.suffix + ".tmp")
    temporary.write_text(json.dumps(instance.to_dict(), indent=2, sort_keys=True), encoding="utf-8")
    temporary.replace(path)


def read_instance(path: Path) -> MarsupialInstance:
    return MarsupialInstance.from_dict(json.loads(Path(path).read_text(encoding="utf-8")))


def write_instance_excel(instance: MarsupialInstance, path: Path) -> None:
    """Write the frozen instance in the legacy centralized-solvers format."""
    from openpyxl import Workbook

    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    workbook.properties.creator = "MRS experiment protocol v1"
    workbook.properties.created = datetime(2000, 1, 1)
    workbook.properties.modified = datetime(2000, 1, 1)
    tasks_sheet = workbook.active
    tasks_sheet.title = "Tasks"
    tasks_sheet.append([
        "task_index",
        "source_x",
        "source_y",
        "destination_x",
        "destination_y",
        "t_delivery",
        "t_operation",
        "t_picking",
    ])
    for task_index, task in enumerate(instance.tasks, start=1):
        tasks_sheet.append([
            task_index,
            task.source[0],
            task.source[1],
            task.destination[0],
            task.destination[1],
            task.due_time,
            task.handling_time,
            task.pickup_time,
        ])

    vehicles_sheet = workbook.create_sheet("Vehicles")
    vehicles_sheet.append([
        "role", "count", "tau_a", "tau_d", "tau_p", "v", "objective_speed"
    ])
    vehicles_sheet.append([
        "MBR",
        instance.n_mbr,
        instance.dock_time,
        instance.detach_time,
        instance.tasks[0].pickup_time,
        instance.resolved_mbr_speed,
        instance.speed,
    ])
    vehicles_sheet.append([
        "DOR", instance.n_dor, None, None, None,
        instance.resolved_dor_speed, instance.speed,
    ])
    positions_sheet = workbook.create_sheet("InitialPositions")
    positions_sheet.append(["role", "robot_id", "x", "y"])
    for robot_id, position in enumerate(instance.mbr_initial_positions):
        positions_sheet.append(["MBR", robot_id, position[0], position[1]])
    for robot_id, position in enumerate(instance.dor_initial_positions):
        positions_sheet.append(["DOR", robot_id, position[0], position[1]])
    temporary = path.with_suffix(path.suffix + ".tmp")
    workbook.save(temporary)
    temporary.replace(path)


def write_instances(instances: Sequence[MarsupialInstance], directory: Path) -> None:
    directory = Path(directory)
    for instance in instances:
        write_instance(instance, directory / f"{instance.instance_id}.json")


def instances_to_tensor_batch(instances: Sequence[MarsupialInstance], device=None):
    """Convert equal-sized physical instances to the centralized model input."""
    import torch

    if not instances:
        raise ValueError("at least one instance is required")
    reference = instances[0]
    signature = (len(reference.tasks), reference.n_mbr, reference.n_dor)
    if any((len(item.tasks), item.n_mbr, item.n_dor) != signature for item in instances):
        raise ValueError("centralized tensor batches require equal task and fleet sizes")
    if any(
        (
            item.speed,
            item.resolved_mbr_speed,
            item.resolved_dor_speed,
            item.dock_time,
            item.detach_time,
        )
        != (
            reference.speed,
            reference.resolved_mbr_speed,
            reference.resolved_dor_speed,
            reference.dock_time,
            reference.detach_time,
        )
        for item in instances
    ):
        raise ValueError("centralized tensor batches require equal physical parameters")
    device = device or torch.device("cpu")
    return {
        "sor_loc": torch.tensor(
            [[task.source for task in item.tasks] for item in instances],
            dtype=torch.float32,
            device=device,
        ),
        "tar_loc": torch.tensor(
            [[task.destination for task in item.tasks] for item in instances],
            dtype=torch.float32,
            device=device,
        ),
        "designated_time": torch.tensor(
            [[[task.due_time] for task in item.tasks] for item in instances],
            dtype=torch.float32,
            device=device,
        ),
        "tau_h": torch.tensor(
            [[[task.handling_time] for task in item.tasks] for item in instances],
            dtype=torch.float32,
            device=device,
        ),
        "tau_p": torch.tensor(
            [[[task.pickup_time] for task in item.tasks] for item in instances],
            dtype=torch.float32,
            device=device,
        ),
        "params": torch.tensor(
            [[item.dock_time, item.detach_time, item.tasks[0].pickup_time] for item in instances],
            dtype=torch.float32,
            device=device,
        ),
        "v": torch.tensor(
            [[item.speed] for item in instances],
            dtype=torch.float32,
            device=device,
        ),
        "v_mbr": torch.tensor(
            [[item.resolved_mbr_speed] for item in instances],
            dtype=torch.float32,
            device=device,
        ),
        "v_dor": torch.tensor(
            [[item.resolved_dor_speed] for item in instances],
            dtype=torch.float32,
            device=device,
        ),
        "mbr_initial_positions": torch.tensor(
            [item.mbr_initial_positions for item in instances],
            dtype=torch.float32,
            device=device,
        ),
        "dor_initial_positions": torch.tensor(
            [item.dor_initial_positions for item in instances],
            dtype=torch.float32,
            device=device,
        ),
        "mom_size": torch.tensor(
            [item.n_mbr for item in instances], dtype=torch.long, device=device
        ),
        "sub_size": torch.tensor(
            [item.n_dor for item in instances], dtype=torch.long, device=device
        ),
    }
